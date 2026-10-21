"""Derived local evidence index; never writes to archive sources or providers.

Only trusted Archive.records() metadata is accepted. Text and paths are returned by
records() for LOCAL consumers; process()/status() contain only safe counts.
"""
from contextlib import contextmanager
import hashlib
import json
from pathlib import Path
import re
import sqlite3
import unicodedata
from io import BytesIO
from pathlib import PurePosixPath
from urllib.parse import quote
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from html.parser import HTMLParser


class _VisibleHTML(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.hidden = []
        self.chunks = []

    def handle_starttag(self, tag, attrs):
        if tag in ('script', 'style', 'head'):
            self.hidden.append(tag)
        elif tag in ('p', 'br', 'div', 'li', 'tr') and not self.hidden:
            self.chunks.append('\n')

    def handle_endtag(self, tag):
        if tag in self.hidden:
            # Malformed markup fails closed: do not release other hidden tags.
            self.hidden.remove(tag)
        elif tag in ('p', 'div', 'li', 'tr') and not self.hidden:
            self.chunks.append('\n')

    def handle_data(self, data):
        if not self.hidden:
            self.chunks.append(data)


EXTRACTION_VERSION = '1'
MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_TEXT_CHARS = 250_000
MAX_PDF_PAGES = 2_000
MAX_ZIP_BYTES = 100 * 1024 * 1024
MAX_ZIP_ENTRIES = 10_000
MAX_PARTS = 20_000
MAX_SHEET_CELLS = 1_000_000


class _LimitExceeded(Exception):
    pass


def _office_zip(data):
    """Validate package before parsers: no extraction, entities, or external IO."""
    with ZipFile(BytesIO(data)) as package:
        infos = package.infolist()
        if len(infos) > MAX_ZIP_ENTRIES:
            raise _LimitExceeded('zip_entries_limit')
        total = 0
        names = set()
        for info in infos:
            name = info.filename
            path = PurePosixPath(name)
            if (path.is_absolute() or '..' in path.parts or '\\' in name or ':' in name
                    or name in names or '\x00' in name):
                raise ValueError('unsafe_zip')
            names.add(name)
            total += info.file_size
            if total > MAX_ZIP_BYTES:
                raise _LimitExceeded('zip_expansion_limit')
            if info.flag_bits & 1:
                raise ValueError('encrypted_zip')
        # OOXML parsers must never see entity/DTD declarations.
        for info in infos:
            if info.filename.endswith(('.xml', '.rels')):
                xml = package.read(info)
                probe = xml.replace(b'\x00', b'').upper()
                if b'<!DOCTYPE' in probe or b'<!ENTITY' in probe:
                    raise ValueError('unsafe_xml')
        return package.read('word/document.xml') if 'word/document.xml' in names else None


class DocumentIndex:
    def __init__(self, root: Path):
        self.root = Path(root).resolve()
        self.root.mkdir(parents=True, exist_ok=True)
        self.database = self.root / 'analysis_documents.sqlite3'
        with self._connect() as db:
            db.execute('CREATE TABLE IF NOT EXISTS documents (document_key TEXT PRIMARY KEY, record TEXT NOT NULL)')
            db.execute('CREATE TABLE IF NOT EXISTS links (message_key TEXT, attachment_index INTEGER, document_key TEXT, record TEXT NOT NULL, PRIMARY KEY(message_key, attachment_index, document_key))')

    @contextmanager
    def _connect(self):
        db = sqlite3.connect(self.database, timeout=30)
        try:
            db.execute('PRAGMA synchronous=FULL')
            with db:
                yield db
        finally:
            db.close()

    def _verify(self, attachment):
        try:
            path = Path(attachment['path'])
            if not path.is_absolute():
                path = self.root / path
            path = path.resolve()
            if not path.is_relative_to(self.root / 'messages'):
                return None, 'failed', 'unsafe_path'
            if not path.is_file():
                return None, 'file_missing', 'file_missing'
            if path.stat().st_size >= MAX_FILE_BYTES:
                return None, 'oversized', 'file_size_limit'
            # Parse these verified bytes, not a second read susceptible to replacement.
            with path.open('rb') as source:
                data = source.read(MAX_FILE_BYTES)
            if len(data) >= MAX_FILE_BYTES:
                return None, 'oversized', 'file_size_limit'
            if hashlib.sha256(data).hexdigest() != attachment.get('sha256', '').lower():
                return None, 'hash_mismatch', 'hash_mismatch'
            return data, 'verified', None
        except (OSError, ValueError, KeyError, TypeError):
            return None, 'failed', 'source_unreadable'

    def _append(self, record, part):
        remaining = MAX_TEXT_CHARS - sum(len(p['text']) for p in record['parts'])
        if len(record['parts']) >= MAX_PARTS or remaining <= 0:
            record['truncated'] = record['review_required'] = True
            return False
        if len(part['text']) > remaining:
            part['text'] = part['text'][:remaining]
            record['truncated'] = record['review_required'] = True
        record['parts'].append(part)
        return not record['truncated']

    def _extract(self, data, record):
        if record['type'] in ('docx', 'xlsx'):
            xml = _office_zip(data)
            prefix = f"attachment:{record['sha256']}"
            if record['type'] == 'docx':
                if xml is None:
                    raise ValueError('missing_document_xml')
                root = ET.fromstring(xml)
                ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
                body = root.find(ns + 'body')
                if body is None:
                    raise ValueError('missing_body')
                tables = {node: i for i, node in enumerate(body.iter(ns + 'tbl'), 1)}
                cells = {}
                for table, ti in tables.items():
                    for ri, row in enumerate(table.findall(ns + 'tr'), 1):
                        for ci, cell in enumerate(row.findall(ns + 'tc'), 1):
                            for p in cell.iter(ns + 'p'):
                                cells[p] = dict(table=ti, row=ri, cell=f'R{ri}C{ci}')
                # Final-view paragraphs; exclude deleted and moved-away revisions.
                def paragraphs(node):
                    if node.tag in (ns + 'del', ns + 'moveFrom'):
                        return
                    if node.tag == ns + 'p':
                        yield node
                    else:
                        for child in node:
                            yield from paragraphs(child)
                def visible(node):
                    if node.tag in (ns + 'del', ns + 'moveFrom'):
                        return ''
                    if node.tag == ns + 't':
                        return node.text or ''
                    if node.tag == ns + 'tab':
                        return '\t'
                    if node.tag in (ns + 'br', ns + 'cr'):
                        return '\n'
                    return ''.join(visible(child) for child in node)
                for i, paragraph in enumerate(paragraphs(body), 1):
                    if not self._append(record, dict(source_id=f'{prefix}:paragraph:{i}', paragraph=i,
                                                     text=visible(paragraph), **cells.get(paragraph, {}))):
                        break
            else:
                from openpyxl import load_workbook
                wb = load_workbook(BytesIO(data), read_only=True, data_only=True, keep_links=False)
                try:
                    visited = 0
                    stop = False
                    for sheet in wb:
                        if sheet.max_row and sheet.max_column and sheet.max_row * sheet.max_column > MAX_SHEET_CELLS:
                            raise _LimitExceeded('sheet_dimensions_limit')
                        for row in sheet.iter_rows():
                            for cell in row:
                                visited += 1
                                if visited > MAX_SHEET_CELLS:
                                    raise _LimitExceeded('cell_visit_limit')
                                if cell.value is not None:
                                    stop = not self._append(record, dict(source_id=f'{prefix}:sheet:{quote(sheet.title, safe="")}:cell:{cell.coordinate}',
                                                                       sheet=sheet.title, row=cell.row, cell=cell.coordinate, text=str(cell.value)))
                                    if stop:
                                        break
                            if stop:
                                break
                        if stop:
                            break
                finally:
                    wb.close()
                record['formula_policy'] = 'cached_values_only; formulas_not_evaluated'
            record['status'] = 'extracted'
            return
        if record['type'] == 'pdf':
            import pymupdf
            with pymupdf.open(stream=data, filetype='pdf') as doc:
                if doc.needs_pass:
                    record.update(status='encrypted', review_required=True)
                    return
                record['page_count'] = len(doc)
                if len(doc) > MAX_PDF_PAGES:
                    record.update(status='oversized', error_code='page_limit', review_required=True)
                    return
                record['ocr_pages'] = []
                remaining = MAX_TEXT_CHARS
                for i, page in enumerate(doc, 1):
                    text = page.get_text('text')
                    if not text.strip():
                        record['ocr_pages'].append(i)
                    if len(text) > remaining:
                        text = text[:remaining]
                        record['truncated'] = record['review_required'] = True
                    remaining -= len(text)
                    record['parts'].append(dict(source_id=f"attachment:{record['sha256']}:page:{i}", page=i, text=text))
                    if remaining == 0 and i < len(doc):
                        record['truncated'] = record['review_required'] = True
                        record['unexamined_pages_from'] = i + 1
                        break
                record['status'] = 'needs_ocr' if record['ocr_pages'] else 'extracted'
                record['review_required'] |= bool(record['ocr_pages'])
            return
        if record['type'] not in ('txt','csv','html','htm','ics'):
            record.update(status='unsupported',review_required=True)
            return
        encoding='utf-16' if data.startswith((b'\xff\xfe',b'\xfe\xff')) else 'utf-8-sig'
        text=data.decode(encoding,errors='replace')
        if record['type'] in ('html','htm'):
            visible=_VisibleHTML();visible.feed(text);visible.close();text=''.join(visible.chunks)
        for i, line in enumerate(text.splitlines(), 1):
            record['parts'].append(dict(source_id=f"attachment:{record['sha256']}:line:{i}", line=i, text=line))
        record['status'] = 'extracted'

    def _bound(self, record):
        remaining = MAX_TEXT_CHARS
        bounded = []
        for part in record['parts']:
            if len(part['text']) > remaining:
                part['text'] = part['text'][:remaining]
                record['truncated'] = record['review_required'] = True
            bounded.append(part)
            remaining -= len(part['text'])
            if remaining == 0:
                if len(bounded) < len(record['parts']):
                    record['truncated'] = record['review_required'] = True
                break
        record['parts'] = bounded

    def process(self, records: list[dict], max_documents: int = 200) -> dict:
        if not isinstance(max_documents, int) or max_documents < 0:
            raise ValueError('max_documents must be a nonnegative integer')
        processed = deferred = 0
        for message in records:
            for position, attachment in enumerate(message.get('attachments', [])):
                sha = str(attachment.get('sha256', '')).lower()
                # Invalid digest gets an opaque, deterministic identity, never a filename.
                valid_sha = bool(re.fullmatch('[0-9a-f]{64}', sha))
                if not valid_sha:
                    sha = hashlib.sha256(json.dumps([message.get('message_key'), position, sha]).encode()).hexdigest()
                key = f'{sha}:{EXTRACTION_VERSION}'
                with self._connect() as db:
                    old = db.execute('SELECT record FROM documents WHERE document_key=?', (key,)).fetchone()
                    if not old and processed >= max_documents:
                        deferred += 1
                        continue
                    data, status, error = self._verify(attachment) if valid_sha else (None, 'failed', 'invalid_sha256')
                    if not old:
                        name = unicodedata.normalize('NFKC', str(attachment.get('original_filename', 'attachment'))).replace('\\', '/').split('/')[-1]
                        name = ''.join(c for c in name if c.isprintable())[:255] or 'attachment'
                        record = dict(document_key=key, sha256=sha, extraction_version=EXTRACTION_VERSION,
                                      type=Path(name).suffix.lower().lstrip('.') or 'unknown', status=status,
                                      path=str(attachment.get('path', '')), filename=name, normalized_filename=name,
                                      parts=[], truncated=False, review_required=False)
                        if error:
                            record['error_code'] = error
                        if data is not None:
                            try:
                                self._extract(data, record)
                                self._bound(record)
                            except _LimitExceeded as exc:
                                record.update(status='oversized', error_code=str(exc), parts=[], review_required=True)
                            except (ImportError,ModuleNotFoundError):
                                record.update(status='missing_dependency',error_code='missing_dependency',parts=[],review_required=True)
                            except Exception:
                                record.update(status='failed', error_code='parse_failed', parts=[], review_required=True)
                        db.execute('INSERT INTO documents VALUES (?,?)', (key, json.dumps(record, ensure_ascii=False)))
                        processed += 1
                    link = dict(message_key=message['message_key'], attachment_index=position,
                                account_id=message.get('account_id'), path=str(attachment.get('path', '')),
                                status=status, error_code=error)
                    db.execute('INSERT OR REPLACE INTO links VALUES (?,?,?,?)',
                               (message['message_key'], position, key, json.dumps(link)))
        return dict(self.status(), processed=processed, deferred=deferred)

    def records(self) -> list[dict]:
        with self._connect() as db:
            result = []
            for key, raw in db.execute('SELECT document_key, record FROM documents ORDER BY document_key'):
                record = json.loads(raw)
                links = [json.loads(row[0]) for row in db.execute('SELECT record FROM links WHERE document_key=? ORDER BY message_key, attachment_index', (key,))]
                record['links'] = links
                record['linked_message_keys'] = sorted({link['message_key'] for link in links})
                result.append(record)
            return result

    def status(self) -> dict:
        result = self.records()
        counts = {}
        for record in result:
            counts[record['status']] = counts.get(record['status'], 0) + 1
        return dict(total=len(result), statuses=counts, ocr_queue=counts.get('needs_ocr', 0),
                    dependency_queue=counts.get('missing_dependency',0),
                    review_required=sum(bool(r.get('review_required')) for r in result))
