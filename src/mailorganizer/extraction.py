"""Bounded offline extraction: no network, macros, formula evaluation or OCR."""
from pathlib import Path
import csv
import io
import re
import zipfile
from xml.etree import ElementTree
from html.parser import HTMLParser

EXTRACTOR_VERSION = 'offline-3'


class UnsafeDocument(ValueError):
    pass


def _result(parts=(), *, status='extracted', complete=True, reason='', source_count=0, review_sources=()):
    parts = list(parts)
    return dict(status=status, text='\n'.join(p['text'] for p in parts), parts=parts,
                complete=complete, requires_review=not complete, reason=reason,
                source_count=source_count, review_sources=list(review_sources))


def _zip_check(path, budget):
    """Check declared AND streamed sizes before a library parses the package."""
    with zipfile.ZipFile(path) as z:
        infos = z.infolist()
        if len(infos) > 10000 or sum(i.file_size for i in infos) > budget:
            raise UnsafeDocument('zip_total_limit')
        total, external = 0, False
        if len({i.filename for i in infos}) != len(infos):
            raise UnsafeDocument('zip_duplicate_members')
        for info in infos:
            name = info.filename.replace('\\', '/')
            if name.startswith('/') or ':' in name or '..' in name.split('/'):
                raise UnsafeDocument('zip_unsafe_path')
            if re.search(r'vba|activex|embeddings/', name, re.I) or Path(name).suffix.lower() in ('.exe', '.dll', '.com', '.bat', '.cmd', '.ps1', '.js', '.vbs'):
                raise UnsafeDocument('zip_executable_or_macro')
            if info.flag_bits & 1 or info.file_size > min(budget, 20000000):
                raise UnsafeDocument('zip_member_limit_or_encrypted')
            count, previous = 0, b''
            with z.open(info) as stream:
                while chunk := stream.read(65536):
                    count += len(chunk)
                    total += len(chunk)
                    if count > min(budget, 20000000) or total > budget:
                        raise UnsafeDocument('zip_actual_limit')
                    if name.lower().endswith(('.xml', '.rels')):
                        scan = (previous + chunk).replace(b'\x00', b'').upper()
                        if b'<!DOCTYPE' in scan or b'<!ENTITY' in scan:
                            raise UnsafeDocument('unsafe_xml_dtd')
                        if b'MACROENABLED' in scan or b'VBAPROJECT' in scan:
                            raise UnsafeDocument('zip_executable_or_macro')
                        previous = chunk[-64:]
            if count != info.file_size:
                raise UnsafeDocument('zip_size_mismatch')
            if name.lower().endswith('.rels'):
                xml = ElementTree.fromstring(z.read(info))
                external |= any(e.get('TargetMode', '').lower() == 'external' for e in xml)
        return external


class _HTML(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.hidden = 0
        self.text = []

    def handle_starttag(self, tag, attrs):
        if tag in ('script', 'style', 'head', 'template'):
            self.hidden += 1
        elif tag in ('p', 'br', 'div', 'li', 'tr') and not self.hidden:
            self.text.append('\n')

    def handle_endtag(self, tag):
        if tag in ('script', 'style', 'head', 'template'):
            self.hidden = max(0, self.hidden - 1)
        elif tag in ('p', 'div', 'li', 'tr') and not self.hidden:
            self.text.append('\n')

    def handle_data(self, data):
        if not self.hidden:
            self.text.append(data)


def _decode(data, mime):
    if data.startswith((b'\xff\xfe', b'\xfe\xff')):
        return data.decode('utf-16')
    charset = re.search(r'charset=["\']?([\w-]+)', mime, re.I)
    return data.decode(charset.group(1) if charset else 'utf-8-sig')


def _bounded(parts, max_chars, *, source_count=None, reasons=(), review_sources=()):
    result, used = [], 0
    reasons = list(reasons)
    for part in parts:
        room = max_chars - used - bool(result)
        if room <= 0:
            reasons.append('max_chars')
            break
        text = part['text'][:room]
        result.append(dict(part, text=text))
        used += len(text) + (1 if len(result) > 1 else 0)
        if len(text) < len(part['text']):
            reasons.append('max_chars')
            break
    reason = ';'.join(dict.fromkeys(reasons))
    return _result(result, status='partial' if reason else 'extracted', complete=not reason,
                   reason=reason, source_count=len(parts) if source_count is None else source_count,
                   review_sources=review_sources)


def extract_document(path: Path, mime: str = '', *, max_pages=200, max_chars=60000, max_bytes=50000000) -> dict:
    """Extract real provenance; unknown/bounded/scan results always require review."""
    if min(max_pages, max_chars, max_bytes) < 1:
        raise ValueError('Extraction limits must be positive')
    path = Path(path)
    try:
        if path.is_symlink() or not path.is_file():
            return _result(status='unsafe', complete=False, reason='not_regular_file')
        if path.stat().st_size > max_bytes:
            return _result(status='too_large', complete=False, reason='max_bytes')
        suffix = path.suffix.lower()
        kind = mime.split(';')[0].strip().lower()
        if suffix == '.pdf' or kind == 'application/pdf':
            import pymupdf
            with pymupdf.open(path) as doc:
                if doc.needs_pass:
                    return _result(status='encrypted', complete=False, reason='password_required', source_count=len(doc))
                parts, review = [], []
                for i in range(min(len(doc), max_pages)):
                    page = doc[i]
                    text = page.get_text()
                    if not text.strip() or page.get_images():
                        review.append(f'page:{i + 1}')
                    parts.append(dict(source_id=f'page:{i + 1}', text=text))
                reasons = (['max_pages'] if len(doc) > max_pages else []) + (['ocr_required_or_image_review'] if review else [])
                return _bounded(parts, max_chars, source_count=len(doc), reasons=reasons, review_sources=review)
        if suffix in ('.docm', '.xlsm', '.exe', '.bat', '.ps1'):
            return _result(status='unsafe', complete=False, reason='executable_or_macro')
        if suffix in ('.docx', '.xlsx') or kind in ('application/vnd.openxmlformats-officedocument.wordprocessingml.document', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'):
            external = _zip_check(path, max_bytes)
            package_reasons = ['external_reference_not_fetched'] if external else []
            if suffix == '.docx' or 'wordprocessingml' in kind:
                from docx import Document
                from docx.text.paragraph import Paragraph
                from docx.table import Table
                doc = Document(path)
                parts, pi, ri = [], 0, 0
                for block in doc.iter_inner_content():
                    if isinstance(block, Paragraph):
                        pi += 1
                        parts.append(dict(source_id=f'paragraph:{pi}', text=block.text))
                    elif isinstance(block, Table):
                        for row in block.rows:
                            ri += 1
                            parts.append(dict(source_id=f'row:{ri}', text='\t'.join(c.text for c in row.cells)))
                # Headers, comments, drawings and tracked changes need explicit review.
                xml = doc.element.xml
                complex_content = any(token in xml for token in ('<w:ins', '<w:del', '<w:drawing', '<w:comment', '<w:txbxContent', '<w:altChunk'))
                with zipfile.ZipFile(path) as z:
                    extra = any(re.match(r'word/(header|footer|footnotes|endnotes|comments)', n) for n in z.namelist())
                return _bounded(parts, max_chars, reasons=package_reasons + (['docx_non_body_or_complex_content'] if extra or complex_content else []))
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            parts, formulas, row_id = [], False, 0
            try:
                for sheet in workbook:
                    for row in sheet.iter_rows():
                        row_id += 1
                        if row_id > 100000:
                            return _bounded(parts, max_chars, reasons=['max_rows'])
                        values = []
                        for cell in row:
                            if cell.data_type == 'f':
                                values.append('[formula omitted]')
                                formulas = True
                            else:
                                values.append('' if cell.value is None else str(cell.value))
                        parts.append(dict(source_id=f'row:{row_id}', sheet=sheet.title, text='\t'.join(values)))
            finally:
                workbook.close()
            return _bounded(parts, max_chars, reasons=package_reasons + (['formula_not_evaluated'] if formulas else []))
        if suffix in ('.txt', '.csv', '.html', '.htm', '.ics') or kind in ('text/plain', 'text/csv', 'text/html', 'text/calendar'):
            with path.open('rb') as stream:
                data = stream.read(max_bytes + 1)
            if len(data) > max_bytes:
                return _result(status='too_large', complete=False, reason='max_bytes')
            text = _decode(data, mime)
            if suffix in ('.html', '.htm') or kind == 'text/html':
                parser = _HTML()
                parser.feed(text)
                text = ''.join(parser.text).strip()
            if suffix == '.csv' or kind == 'text/csv':
                parts = [dict(source_id=f'row:{i}', text='\t'.join(row)) for i, row in enumerate(csv.reader(io.StringIO(text)), 1)]
            else:
                parts = [dict(source_id=f'paragraph:{i}', text=line) for i, line in enumerate(text.splitlines(), 1)]
            return _bounded(parts, max_chars)
        return _result(status='unsupported', complete=False, reason='unsupported_format')
    except UnsafeDocument as exc:
        return _result(status='unsafe', complete=False, reason=str(exc))
    except ImportError:
        return _result(status='dependency_missing', complete=False, reason='optional_local_dependency_missing')
    except (UnicodeError, LookupError):
        return _result(status='decode_error', complete=False, reason='charset_required_or_invalid')
    except Exception:
        return _result(status='corrupt', complete=False, reason='unreadable_or_invalid_document')
