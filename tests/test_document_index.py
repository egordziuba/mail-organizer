"""Synthetic fixtures only: archive originals must remain immutable."""
import hashlib
import json
from pathlib import Path


def attachment(root, name, data, message='m1'):
    path = root / 'messages' / message / 'attachments' / name
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    return {'message_key': message, 'account_id': 'synthetic', 'attachments': [
        {'path': str(path), 'original_filename': name, 'sha256': hashlib.sha256(data).hexdigest(), 'size': len(data)}]}


def test_text_persists_evidence_and_deduplicates_without_changing_originals(tmp_path):
    from mailorganizer.document_index import DocumentIndex
    data = 'Private synthetic text\nSecond line'.encode()
    first = attachment(tmp_path, 'note.txt', data)
    second = attachment(tmp_path, 'other.txt', data, 'm2')
    before = json.dumps([first, second], sort_keys=True)
    index = DocumentIndex(tmp_path)
    summary = index.process([first, second])
    assert summary['total'] == 1
    assert summary['processed'] == 1
    assert 'Private' not in json.dumps(summary)
    result = index.records()[0]
    assert result['status'] == 'extracted'
    assert result['linked_message_keys'] == ['m1', 'm2']
    assert result['parts'][0]['source_id'] == f"attachment:{result['sha256']}:line:1"
    assert result['parts'][1]['text'] == 'Second line'
    assert DocumentIndex(tmp_path).process([first, second])['processed'] == 0
    assert json.dumps([first, second], sort_keys=True) == before
    assert Path(first['attachments'][0]['path']).read_bytes() == data


def test_untrusted_paths_hashes_missing_and_limits_are_isolated(tmp_path):
    from mailorganizer.document_index import DocumentIndex
    good = attachment(tmp_path, 'ok.txt', b'ok', 'ok')
    missing = attachment(tmp_path, 'missing.txt', b'missing', 'missing')
    Path(missing['attachments'][0]['path']).unlink()
    mismatch = attachment(tmp_path, 'changed.txt', b'before', 'changed')
    Path(mismatch['attachments'][0]['path']).write_bytes(b'after')
    outside = attachment(tmp_path, 'outside.txt', b'outside', 'outside')
    outside['attachments'][0]['path'] = str(tmp_path / '..' / 'outside.txt')
    index = DocumentIndex(tmp_path)
    index.process([missing, mismatch, outside, good], max_documents=2)
    assert index.status()['total'] == 2
    assert index.status()['statuses'] == {'file_missing': 1, 'hash_mismatch': 1}
    index.process([missing, mismatch, outside, good])
    assert index.status()['total'] == 4
    assert {r['status'] for r in index.records()} == {'file_missing', 'hash_mismatch', 'failed', 'extracted'}
    assert next(r for r in index.records() if r['status'] == 'failed')['error_code'] == 'unsafe_path'


def test_deduplicated_link_must_still_verify_its_own_bytes(tmp_path):
    from mailorganizer.document_index import DocumentIndex
    first = attachment(tmp_path, 'one.txt', b'same')
    second = attachment(tmp_path, 'two.txt', b'same', 'm2')
    Path(second['attachments'][0]['path']).write_bytes(b'changed')
    index = DocumentIndex(tmp_path)
    index.process([first, second])
    record = index.records()[0]
    assert record['links'][1]['status'] == 'hash_mismatch'
    assert record['links'][0]['status'] == 'verified'


def test_size_and_text_bounds_are_explicit(tmp_path, monkeypatch):
    import mailorganizer.document_index as module
    monkeypatch.setattr(module, 'MAX_FILE_BYTES', 10)
    oversized = attachment(tmp_path, 'big.txt', b'12345678901')
    index = module.DocumentIndex(tmp_path)
    index.process([oversized])
    assert index.records()[0]['status'] == 'oversized'
    monkeypatch.setattr(module, 'MAX_FILE_BYTES', 100)
    monkeypatch.setattr(module, 'MAX_TEXT_CHARS', 5)
    small = attachment(tmp_path, 'small.txt', b'abcdefgh', 'm2')
    index.process([small])
    result = next(r for r in index.records() if r['status'] == 'extracted')
    assert result['parts'][0]['text'] == 'abcde'
    assert result['truncated'] and result['review_required']


def test_pdf_pages_scan_queue_encryption_and_page_limit(tmp_path, monkeypatch):
    import pymupdf
    import mailorganizer.document_index as module
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), 'Synthetic first page')
    doc.new_page()  # no OCR must be invented for this page
    raw = doc.tobytes()
    pdf = attachment(tmp_path, 'sample.pdf', raw)
    scan_doc = pymupdf.open()
    scan_doc.new_page()
    scan = attachment(tmp_path, 'scan.pdf', scan_doc.tobytes(), 'scan')
    encrypted = attachment(tmp_path, 'locked.pdf', doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw='owner', user_pw='test'), 'locked')
    index = module.DocumentIndex(tmp_path)
    index.process([pdf, scan, encrypted])
    by_name = {r['filename']: r for r in index.records()}
    result = by_name['sample.pdf']
    assert result['status'] == 'needs_ocr'
    assert result['parts'][0]['page'] == 1
    assert result['parts'][0]['source_id'] == f"attachment:{result['sha256']}:page:1"
    assert 'Synthetic first page' in result['parts'][0]['text']
    assert result['ocr_pages'] == [2]
    assert by_name['scan.pdf']['status'] == 'needs_ocr'
    assert by_name['scan.pdf']['parts'][0]['text'] == ''
    assert by_name['locked.pdf']['status'] == 'encrypted'
    assert index.status()['ocr_queue'] == 2
    assert Path(pdf['attachments'][0]['path']).read_bytes() == raw
    monkeypatch.setattr(module, 'MAX_PDF_PAGES', 1)
    limited = attachment(tmp_path, 'limit.pdf', raw + b'\n', 'limit')
    index.process([limited])
    assert next(r for r in index.records() if r['filename'] == 'limit.pdf')['status'] == 'oversized'
    doc.close()
    scan_doc.close()


def test_docx_paragraph_indices_include_blanks_and_table_cells(tmp_path):
    from io import BytesIO
    from docx import Document
    from mailorganizer.document_index import DocumentIndex
    doc = Document()
    doc.add_paragraph('First')
    doc.add_paragraph('')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Table A'
    table.cell(0, 1).text = 'Table B'
    doc.add_paragraph('Last')
    stream = BytesIO()
    doc.save(stream)
    index = DocumentIndex(tmp_path)
    index.process([attachment(tmp_path, 'sample.docx', stream.getvalue())])
    result = index.records()[0]
    assert result['status'] == 'extracted'
    assert [p['text'] for p in result['parts']] == ['First', '', 'Table A', 'Table B', 'Last']
    assert [p['paragraph'] for p in result['parts']] == [1, 2, 3, 4, 5]
    assert result['parts'][2]['table'] == 1
    assert result['parts'][2]['cell'] == 'R1C1'
    assert all('page' not in p for p in result['parts'])


def test_xlsx_cell_evidence_and_formula_not_evaluated(tmp_path):
    from io import BytesIO
    from openpyxl import Workbook
    from mailorganizer.document_index import DocumentIndex
    wb = Workbook()
    wb.active.title = 'Budget'
    wb.active['A1'] = 'Synthetic label'
    wb.active['B3'] = 42
    wb.active['C3'] = '=HYPERLINK("https://invalid.example/", "do not visit")'
    wb.create_sheet('Next')['A2'] = 'Second sheet'
    stream = BytesIO()
    wb.save(stream)
    index = DocumentIndex(tmp_path)
    index.process([attachment(tmp_path, 'sample.xlsx', stream.getvalue())])
    result = index.records()[0]
    assert result['status'] == 'extracted'
    assert [(p['sheet'], p['cell'], p['text']) for p in result['parts']] == [('Budget', 'A1', 'Synthetic label'), ('Budget', 'B3', '42'), ('Next', 'A2', 'Second sheet')]
    assert result['parts'][1]['row'] == 3
    assert result['parts'][0]['source_id'].endswith(':sheet:Budget:cell:A1')


def test_office_zip_bomb_paths_and_malformed_are_safe(tmp_path, monkeypatch):
    from io import BytesIO
    from zipfile import ZipFile
    import mailorganizer.document_index as module
    index = module.DocumentIndex(tmp_path)
    stream = BytesIO()
    with ZipFile(stream, 'w') as z:
        z.writestr('../escape.xml', 'not a document')
    bad_path = attachment(tmp_path, 'bad.docx', stream.getvalue(), 'bad')
    stream = BytesIO()
    with ZipFile(stream, 'w') as z:
        z.writestr('word/document.xml', 'x' * 1024)
    bomb = attachment(tmp_path, 'bomb.docx', stream.getvalue(), 'bomb')
    monkeypatch.setattr(module, 'MAX_ZIP_BYTES', 100)
    index.process([bad_path, bomb, attachment(tmp_path, 'malformed.xlsx', b'not zip', 'broken')])
    by_name = {r['filename']: r for r in index.records()}
    assert by_name['bad.docx']['status'] == 'failed'
    assert by_name['bomb.docx']['status'] == 'oversized'
    assert by_name['malformed.xlsx']['status'] == 'failed'


def test_html_script_style_removed_text_decode_csv_and_unsupported(tmp_path):
    from mailorganizer.document_index import DocumentIndex
    html = attachment(tmp_path, 'page.html', b'<style>secret style</style><p>Visible &amp; safe</p><script>secret script</script>Tail', 'html')
    csv = attachment(tmp_path, 'data.csv', b'name,value\nExample,12', 'csv')
    utf16 = attachment(tmp_path, 'unicode.txt', 'Test UTF16'.encode('utf-16'), 'utf16')
    image = attachment(tmp_path, 'image.png', b'\x89PNG\r\n', 'image')
    index = DocumentIndex(tmp_path)
    index.process([html, csv, utf16, image])
    by_name = {r['filename']: r for r in index.records()}
    text = '\n'.join(p['text'] for p in by_name['page.html']['parts'])
    assert 'Visible & safe' in text and 'Tail' in text
    assert 'secret' not in text
    assert by_name['data.csv']['parts'][1]['line'] == 2
    assert by_name['unicode.txt']['parts'][0]['text'] == 'Test UTF16'
    assert by_name['image.png']['status'] == 'unsupported'
    assert by_name['image.png']['parts'] == []


def test_missing_dependency_is_queued_not_unsupported(tmp_path, monkeypatch):
    import builtins
    from mailorganizer.document_index import DocumentIndex
    original_import = builtins.__import__
    def blocked(name, *args, **kwargs):
        if name == 'pymupdf':
            raise ModuleNotFoundError('synthetic dependency unavailable')
        return original_import(name, *args, **kwargs)
    monkeypatch.setattr(builtins, '__import__', blocked)
    index = DocumentIndex(tmp_path)
    index.process([attachment(tmp_path, 'doc.pdf', b'%PDF-synthetic')])
    assert index.records()[0]['status'] == 'missing_dependency'
    assert index.records()[0]['error_code'] == 'missing_dependency'
    assert index.status()['dependency_queue'] == 1


def test_unsupported_malformed_digest_and_negative_limit(tmp_path):
    import pytest
    from mailorganizer.document_index import DocumentIndex
    item = attachment(tmp_path, 'bad.txt', b'bad')
    item['attachments'][0]['sha256'] = '../not-a-hash'
    index = DocumentIndex(tmp_path)
    index.process([item])
    assert index.records()[0]['error_code'] == 'invalid_sha256'
    with pytest.raises(ValueError):
        index.process([], max_documents=-1)
