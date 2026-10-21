from pathlib import Path
import hashlib
import pytest
from mailorganizer import extraction


def test_pdf_scan_mixed_encrypted_and_corrupt(tmp_path):
    import fitz
    path = tmp_path / 'scan.pdf'
    with fitz.open() as doc:
        doc.new_page().insert_text((72, 72), 'Visible')
        doc.new_page()
        doc.save(path)
        locked = tmp_path / 'locked.pdf'
        doc.save(locked, encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw='owner', user_pw='secret')
    result = extraction.extract_document(path)
    assert not result['complete'] and result['requires_review']
    assert 'ocr_required' in result['reason'] and result['review_sources'] == ['page:2']
    assert extraction.extract_document(locked)['status'] == 'encrypted'
    bad = tmp_path / 'bad.pdf'
    bad.write_bytes(b'not a pdf')
    assert extraction.extract_document(bad)['status'] == 'corrupt'


def test_docx_structure_and_zip_safety(tmp_path):
    from docx import Document
    import zipfile
    path = tmp_path / 'words.docx'
    doc = Document()
    doc.add_paragraph('Привет Universität')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Table body'
    doc.save(path)
    result = extraction.extract_document(path)
    assert result['complete'] and 'Table body' in result['text']
    assert result['parts'][0]['source_id'] == 'paragraph:1'
    assert all(not p['source_id'].startswith('page:') for p in result['parts'])
    bomb = tmp_path / 'bomb.docx'
    with zipfile.ZipFile(bomb, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('word/document.xml', 'x' * 1000000)
    assert extraction.extract_document(bomb, max_bytes=10000)['status'] == 'unsafe'
    with zipfile.ZipFile(bomb, 'w') as z:
        z.writestr('word/vbaProject.bin', b'no execution')
    assert extraction.extract_document(bomb)['status'] == 'unsafe'
    with zipfile.ZipFile(bomb, 'w') as z:
        z.writestr('word/document.xml', '<!DOCTYPE x SYSTEM "https://invalid.test/evil"><x/>')
    assert extraction.extract_document(bomb)['status'] == 'unsafe'


def test_docx_external_references_never_fetched_and_are_reviewed(tmp_path):
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    path = tmp_path / 'external.docx'
    doc = Document()
    doc.add_paragraph('Only local text')
    doc.part.relate_to('https://invalid.test/private', RT.HYPERLINK, is_external=True)
    doc.save(path)
    result = extraction.extract_document(path)
    assert result['requires_review'] and 'external_reference_not_fetched' in result['reason']
    assert result['text'] == 'Only local text'


def test_text_csv_html_xlsx_and_truncation(tmp_path):
    path = tmp_path / 'text.txt'
    path.write_bytes('Привет'.encode('utf-16'))
    assert extraction.extract_document(path)['text'] == 'Привет'
    path.write_text('abcdefgh', encoding='utf-8')
    result = extraction.extract_document(path, max_chars=3)
    assert result['text'] == 'abc' and not result['complete'] and 'max_chars' in result['reason']
    csv = tmp_path / 'table.csv'
    csv.write_text('name,amount\n"Üser",42', encoding='utf-8')
    assert extraction.extract_document(csv)['parts'][1]['source_id'] == 'row:2'
    html = tmp_path / 'test.html'
    html.write_text('<p>Safe</p><script>evil()</script><style>hide</style>', encoding='utf-8')
    assert extraction.extract_document(html)['text'].strip() == 'Safe'
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.active.append(['Unicode', 'Ж'])
    workbook.active.append(['=HYPERLINK("https://invalid.test")', 7])
    xlsx = tmp_path / 'book.xlsx'
    workbook.save(xlsx)
    result = extraction.extract_document(xlsx)
    assert 'Ж' in result['text'] and result['parts'][0]['source_id'] == 'row:1'
    assert 'https' not in result['text'] and result['requires_review']


def test_pdf_real_pages_and_limits_preserve_original(tmp_path):
    import fitz
    path = tmp_path / 'document.pdf'
    with fitz.open() as doc:
        for text in ('First page', 'Second page'):
            doc.new_page().insert_text((72, 72), text)
        doc.save(path)
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    result = extraction.extract_document(path)
    assert result['status'] == 'extracted'
    assert result['complete']
    assert [p['source_id'] for p in result['parts']] == ['page:1', 'page:2']
    assert 'Second page' in result['text']
    limited = extraction.extract_document(path, max_pages=1)
    assert not limited['complete'] and limited['requires_review']
    assert limited['source_count'] == 2 and 'max_pages' in limited['reason']
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before
